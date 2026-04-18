import sqlite3
from datetime import datetime, timezone, timedelta
from pathlib import Path
from services.db.connection import DB_PATH
from flask import Flask, render_template, request, redirect, url_for, session
import os
import hmac
import secrets
from io import BytesIO
from openpyxl import Workbook
from routes import register_core_routes, register_match_routes, register_planner_routes
from services.ronin_api import fetch_all_owned_chickens
from services.metadata_parser import parse_chicken_record
from services.match_rules import (
    find_potential_matches,
    is_generation_gap_allowed,
    is_parent_offspring,
    is_full_siblings,
)
from services.family_roots import (
    resolve_family_roots_for_all,
    complete_ninuno_via_lineage_with_resume,
    initialize_simple_family_roots_for_wallet,
)
from services.market_candidate_cache import init_market_candidate_cache_table
from services.market_featured_service import get_featured_market_feed
from services.market_listing_cache import init_market_listing_cache_table

from services.chicken_enricher import enrich_chicken_records
from services.gene_build_picker import get_best_available_gene_build_info
from services.build_eval import evaluate_build, count_added_missing_traits
from services.wallet_access import get_wallet_access_expiry_display
from services.gene_classifier import classify_gene_profile
from services.database import (
    init_db,
    upsert_chicken,
    get_chickens_by_wallet,
    clear_family_roots_for_wallet,
    clear_family_roots_for_token,
    clear_stale_family_root_summaries,
    upsert_family_root_summary,
    insert_family_root_items,
    get_static_chickens_by_token_ids,
    preload_cached_family_roots_for_wallet,
    get_wallet_last_synced_at,
    upsert_wallet_last_synced_at,
    delete_wallet_chickens_not_in_tokens,
)
from services.ip_breeding import (
    recommend_ip_item,
    get_ip_item_candidates,
    resolve_pair_item_recommendations,
    get_effective_ip_stat,
    get_weakest_ip_stat_info,
    build_ip_pair_quality,
    pair_has_usable_ip_items,
    normalize_auto_ninuno_filter,
    chicken_passes_auto_ninuno_filter,
    build_ip_available_auto_candidates,
    pick_best_ip_auto_match,
    pick_best_ip_auto_match_from_pool,
    build_ip_multi_matches,
    sort_ip_match_rows,
)

from services.gene_breeding import (
    recommend_gene_item,
    get_gene_item_candidates,
    get_gene_build_target_info,
    get_gene_pair_completion,
    get_gene_pair_completion_from_row,
    build_gene_pair_quality,
    build_gene_potential_matches,
    build_gene_potential_matches_strict,
    build_gene_available_auto_candidates_same_build,
    pick_best_gene_auto_match,
    pick_best_gene_auto_match_from_pool,
    normalize_instinct_name,
    get_instinct_tier_rank,
    build_prefers_instinct,
    sort_gene_match_rows,

)

from services.ultimate_breeding import (
    is_ultimate_eligible,
    get_ultimate_type_display,
    get_ultimate_build_display,
    filter_and_sort_ultimate_candidates,
    get_ultimate_item_candidates,
    resolve_ultimate_pair_item_recommendations,
    build_ultimate_pair_quality_from_items,
    pick_best_ultimate_auto_match as service_pick_best_ultimate_auto_match,
    build_ultimate_available_auto_candidates as service_build_ultimate_available_auto_candidates,
    refresh_ultimate_primary_builds_if_needed,
)

from services.wallet_access import (
    init_wallet_access_db,
    is_valid_wallet,
    has_wallet_access,
    set_authorized_wallet,
    is_authorized_wallet,
    get_wallet_access_expiry_display,
    grant_manual_access,
    get_wallet_access_rows,
    format_wallet_access_rows,
    has_active_payment_access_in_db,
)

from services.ip_available_table import (
    parse_csv_query_values,
    normalize_ip_available_ninuno_filter,
    enrich_ip_available_chicken_row,
    build_ip_available_filter_options,
    chicken_matches_ip_available_filters,
    build_ip_active_filters,
    sort_ip_available_chickens,
)

from services.gene_available_table import (
    GENE_BUILD_ORDER,
    parse_csv_query_values as parse_gene_csv_query_values,
    normalize_gene_available_build_filter,
    normalize_gene_available_ninuno_filter,
    normalize_gene_available_source_values,
    enrich_gene_available_chicken_row,
    build_gene_available_filter_options,
    chicken_matches_gene_available_filters,
    build_gene_active_filters,
    sort_gene_available_chickens,
)

from services.ultimate_available_table import (
    parse_csv_query_values as parse_ultimate_csv_query_values,
    normalize_ultimate_available_ninuno_filter,
    normalize_ultimate_build_value,
    enrich_ultimate_available_chicken_row,
    build_ultimate_available_filter_options,
    chicken_matches_ultimate_available_filters,
    build_ultimate_active_filters,
    sort_ultimate_available_chickens as sort_ultimate_available_table_chickens,
)
from services.planner_item_requirements import (
    build_wallet_planner_item_requirements_summary,
    build_per_pair_item_status,
)
from services.wallet_item_inventory import build_wallet_inventory_lookup
from services.planner_bookmarklet import (
    build_apex_breeder_bookmarklet_code,
    build_bookmarklet_inventory_name_lookup,
)
app = Flask(__name__)


def env_flag(name, default=False):
    value = os.environ.get(name)
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


FLASK_DEBUG_ENABLED = env_flag("FLASK_DEBUG", default=False)
SESSION_COOKIE_SECURE_ENABLED = env_flag(
    "SESSION_COOKIE_SECURE",
    default=bool(os.environ.get("RAILWAY_ENVIRONMENT", "").strip()),
)
session_secret = os.environ.get("FLASK_SECRET_KEY", "").strip()
if not session_secret:
    session_secret = secrets.token_hex(32)

app.secret_key = session_secret
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=SESSION_COOKIE_SECURE_ENABLED,
)

CONTRACTS = [
    "0xee9436518030616bc315665678738a4348463df4",
    "0x322b3d98ddbd589dc2e8dd83659bb069828231e0",
]

MATCH_SETTINGS = {
    "max_generation_gap": 3,
    "max_ip_diff": 10,
    "max_breed_count_diff": 1,
}

init_db()
init_wallet_access_db()
init_market_candidate_cache_table()
init_market_listing_cache_table()

OWNER_ADMIN_PASSWORD = os.environ.get("OWNER_ADMIN_PASSWORD", "").strip()
OWNER_WHITELIST_ROUTE = "/owner/grant-access"
CSRF_SESSION_KEY = "_csrf_token"
CSRF_COOKIE_NAME = "apex_csrf_token"
CSRF_HEADER_NAME = "X-CSRF-Token"
CSRF_FORM_FIELD = "csrf_token"
OWNER_ADMIN_MAX_ATTEMPTS = 5
OWNER_ADMIN_LOCK_MINUTES = 15

STATIC_EXPORT_DB_PATH = Path(__file__).resolve().parent / "cache" / "chicken_static_export.db"


def get_csrf_token():
    token = session.get(CSRF_SESSION_KEY)
    if not token:
        token = secrets.token_urlsafe(32)
        session[CSRF_SESSION_KEY] = token
    return token


def is_ajax_request():
    requested_with = str(request.headers.get("X-Requested-With") or "").strip().lower()
    accepts_json = "application/json" in str(request.headers.get("Accept") or "").lower()
    return requested_with == "xmlhttprequest" or accepts_json


def build_csrf_error_response():
    if is_ajax_request():
        return {"ok": False, "error": "CSRF validation failed."}, 400
    return "CSRF validation failed.", 400


def validate_csrf_request():
    expected_token = str(session.get(CSRF_SESSION_KEY) or "").strip()
    submitted_token = str(
        request.form.get(CSRF_FORM_FIELD)
        or request.headers.get(CSRF_HEADER_NAME)
        or request.headers.get("X-CSRFToken")
        or ""
    ).strip()
    cookie_token = str(request.cookies.get(CSRF_COOKIE_NAME) or "").strip()

    if not expected_token or not submitted_token or not cookie_token:
        return False

    return (
        hmac.compare_digest(expected_token, submitted_token)
        and hmac.compare_digest(expected_token, cookie_token)
    )


def is_owner_admin_locked():
    locked_until_raw = session.get("owner_admin_locked_until")
    if not locked_until_raw:
        return False, 0

    try:
        locked_until = datetime.fromisoformat(locked_until_raw)
    except Exception:
        session.pop("owner_admin_locked_until", None)
        return False, 0

    now = datetime.now(timezone.utc)
    if locked_until <= now:
        session.pop("owner_admin_locked_until", None)
        session.pop("owner_admin_failed_attempts", None)
        return False, 0

    remaining_minutes = max(1, int((locked_until - now).total_seconds() // 60) + 1)
    return True, remaining_minutes


def register_owner_admin_failure():
    failures = int(session.get("owner_admin_failed_attempts", 0)) + 1
    session["owner_admin_failed_attempts"] = failures

    if failures >= OWNER_ADMIN_MAX_ATTEMPTS:
        locked_until = datetime.now(timezone.utc).replace(microsecond=0) + timedelta(minutes=OWNER_ADMIN_LOCK_MINUTES)
        session["owner_admin_locked_until"] = locked_until.isoformat()


def clear_owner_admin_failures():
    session.pop("owner_admin_failed_attempts", None)
    session.pop("owner_admin_locked_until", None)


def owner_password_is_valid(owner_password):
    if not OWNER_ADMIN_PASSWORD:
        return False
    return hmac.compare_digest(owner_password, OWNER_ADMIN_PASSWORD)


@app.context_processor
def inject_csrf_token():
    return {"csrf_token": get_csrf_token()}


@app.before_request
def protect_post_requests():
    get_csrf_token()

    if request.method == "POST" and not validate_csrf_request():
        return build_csrf_error_response()


@app.after_request
def persist_csrf_cookie(response):
    response.set_cookie(
        CSRF_COOKIE_NAME,
        get_csrf_token(),
        secure=app.config.get("SESSION_COOKIE_SECURE", False),
        httponly=False,
        samesite=app.config.get("SESSION_COOKIE_SAMESITE", "Lax"),
    )
    return response

def quote_sqlite_identifier(name):
    return '"' + str(name).replace('"', '""') + '"'

def sync_static_export_tables_to_main_db(source_path=None):
    source_path = Path(source_path or STATIC_EXPORT_DB_PATH)

    if not source_path.exists():
        raise FileNotFoundError(f"Static export DB not found: {source_path}")

    synced = []

    with sqlite3.connect(source_path) as source_conn, sqlite3.connect(DB_PATH) as dest_conn:
        source_conn.row_factory = sqlite3.Row
        dest_conn.row_factory = sqlite3.Row

        table_rows = source_conn.execute(
            """
            SELECT name, sql
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            ORDER BY name
            """
        ).fetchall()

        if not table_rows:
            raise ValueError("No user tables found in the static export DB.")

        for row in table_rows:
            table_name = str(row["name"] or "").strip()
            create_sql = row["sql"]

            if not table_name or not create_sql:
                continue

            quoted_table = quote_sqlite_identifier(table_name)

            dest_conn.execute(f"DROP TABLE IF EXISTS {quoted_table}")
            dest_conn.execute(create_sql)

            source_rows = source_conn.execute(f"SELECT * FROM {quoted_table}").fetchall()

            if source_rows:
                column_names = [col[1] for col in source_conn.execute(f"PRAGMA table_info({quoted_table})").fetchall()]
                quoted_columns = ", ".join(quote_sqlite_identifier(col) for col in column_names)
                placeholders = ", ".join(["?"] * len(column_names))

                dest_conn.executemany(
                    f"INSERT INTO {quoted_table} ({quoted_columns}) VALUES ({placeholders})",
                    [tuple(row[col] for col in column_names) for row in source_rows],
                )

            index_rows = source_conn.execute(
                """
                SELECT sql
                FROM sqlite_master
                WHERE type = 'index'
                  AND tbl_name = ?
                  AND sql IS NOT NULL
                ORDER BY name
                """,
                (table_name,),
            ).fetchall()

            for index_row in index_rows:
                index_sql = index_row["sql"]
                if index_sql:
                    try:
                        dest_conn.execute(index_sql)
                    except sqlite3.OperationalError:
                        pass

            row_count = dest_conn.execute(f"SELECT COUNT(*) FROM {quoted_table}").fetchone()[0]
            synced.append({
                "table": table_name,
                "row_count": row_count,
            })

        dest_conn.commit()

    return synced

def get_required_static_cache_tables():
    return ["chicken_static"]


def get_missing_static_cache_tables():
    required_tables = get_required_static_cache_tables()

    with sqlite3.connect(DB_PATH) as conn:
        existing_rows = conn.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            """
        ).fetchall()

    existing_names = {str(row[0]).strip().lower() for row in existing_rows}
    return [
        table_name
        for table_name in required_tables
        if table_name.strip().lower() not in existing_names
    ]


def ensure_static_cache_tables_loaded():
    missing_tables = get_missing_static_cache_tables()

    if not missing_tables:
        return {
            "loaded": False,
            "synced_tables": [],
            "missing_tables": [],
        }

    if not STATIC_EXPORT_DB_PATH.exists():
        return {
            "loaded": False,
            "synced_tables": [],
            "missing_tables": missing_tables,
        }

    sync_results = sync_static_export_tables_to_main_db()

    return {
        "loaded": True,
        "synced_tables": sync_results,
        "missing_tables": missing_tables,
    }


EGG_STATIC_CACHE_FIELDS = [
    "parent_1",
    "parent_2",
]

HATCHED_STATIC_CACHE_FIELDS = [
    "parent_1",
    "parent_2",
    "gender",
    "generation_text",
    "generation_num",
    "instinct",
    "beak",
    "comb",
    "eyes",
    "feet",
    "wings",
    "tail",
    "body",
    "beak_h1",
    "beak_h2",
    "beak_h3",
    "comb_h1",
    "comb_h2",
    "comb_h3",
    "eyes_h1",
    "eyes_h2",
    "eyes_h3",
    "feet_h1",
    "feet_h2",
    "feet_h3",
    "wings_h1",
    "wings_h2",
    "wings_h3",
    "tail_h1",
    "tail_h2",
    "tail_h3",
    "body_h1",
    "body_h2",
    "body_h3",
    "innate_attack",
    "innate_defense",
    "innate_speed",
    "innate_health",
    "innate_ferocity",
    "innate_cockrage",
    "innate_evasion",
]

def merge_static_chicken_cache(record, static_row):
    if not static_row:
        return record

    is_live_egg = bool(record.get("is_egg")) or str(record.get("type") or "").strip().lower() == "egg"
    allowed_fields = EGG_STATIC_CACHE_FIELDS if is_live_egg else HATCHED_STATIC_CACHE_FIELDS

    for field in allowed_fields:
        if field not in static_row:
            continue

        current_value = record.get(field)
        if current_value not in (None, ""):
            continue

        static_value = static_row.get(field)
        if static_value is None or static_value == "":
            continue

        record[field] = static_value

    return record

def needs_recessive_enrichment(chicken):
    return not chicken.get("gene_profile_loaded")


def enrich_missing_recessive_data_in_batches(chickens, wallet, page_key, batch_size=5, prioritized_token_id=None):
    missing = [row for row in chickens if is_breedable(row) and needs_recessive_enrichment(row)]

    if not missing:
        session[f"{page_key}_cursor_{wallet}"] = 0
        return {
            "loaded": 0,
            "remaining": 0,
        }

    prioritized = []
    remaining = missing

    if prioritized_token_id:
        prioritized = [
            row for row in missing
            if str(row.get("token_id") or "") == str(prioritized_token_id)
        ]
        remaining = [
            row for row in missing
            if str(row.get("token_id") or "") != str(prioritized_token_id)
        ]

    cursor_key = f"{page_key}_cursor_{wallet}"
    cursor = safe_int(session.get(cursor_key), 0)
    if cursor is None:
        cursor = 0

    if remaining:
        if cursor >= len(remaining):
            cursor = 0
        rotated = remaining[cursor:] + remaining[:cursor]
    else:
        rotated = []

    selected_batch = prioritized[:1]
    remaining_slots = max(0, batch_size - len(selected_batch))
    selected_batch.extend(rotated[:remaining_slots])

    if not selected_batch:
        return {
            "loaded": 0,
            "remaining": len(missing),
        }

    enriched = enrich_chicken_records(selected_batch)
    loaded_count = 0

    for chicken in enriched:
        refreshed = apply_gene_profile_classification(chicken)
        upsert_chicken(refreshed)
        if refreshed.get("gene_profile_loaded"):
            loaded_count += 1
            
    if remaining:
        next_cursor = cursor + remaining_slots
        if next_cursor >= len(remaining):
            next_cursor = 0
        session[cursor_key] = next_cursor
    else:
        session[cursor_key] = 0

    refreshed = get_chickens_by_wallet(wallet)
    remaining_after = sum(1 for row in refreshed if is_breedable(row) and needs_recessive_enrichment(row))

    return {
        "loaded": loaded_count,
        "remaining": remaining_after,
    }

def has_cached_recessive_ready_data(record):
    return bool(
        record.get("gene_profile_loaded")
        or record.get("recessive_build")
        or any(
            str(record.get(field) or "").strip()
            for field in [
                "beak_h1", "beak_h2", "beak_h3",
                "comb_h1", "comb_h2", "comb_h3",
                "eyes_h1", "eyes_h2", "eyes_h3",
                "feet_h1", "feet_h2", "feet_h3",
                "wings_h1", "wings_h2", "wings_h3",
                "tail_h1", "tail_h2", "tail_h3",
                "body_h1", "body_h2", "body_h3",
            ]
        )
    )

ITEM_IMAGE_URLS = {
    "Soulknot": "https://app.chickensaga.com/images/crafting/SOULKNOT.webp",
    "Gregor's Gift": "https://app.chickensaga.com/images/crafting/GREGOR%27S%20GIFT.webp",
    "Mendel's Memento": "https://app.chickensaga.com/images/crafting/MENDEL%27S%20MEMENTO.webp",
    "Quentin's Talon": "https://app.chickensaga.com/images/crafting/QUENTIN%27S%20TALON.webp",
    "Dragon's Whip": "https://app.chickensaga.com/images/crafting/DRAGON%27S%20WHIP.webp",
    "Chibidei's Curse": "https://app.chickensaga.com/images/crafting/CHIBIDEI%27S%20CURSE.webp",
    "All-seeing Seed": "https://app.chickensaga.com/images/crafting/ALL-SEEING%20SEED.webp",
    "Chim Lac's Curio": "https://app.chickensaga.com/images/crafting/CHIM%20LAC%27S%20CURIO.webp",
    "Suave Scissors": "https://app.chickensaga.com/images/crafting/SUAVE%20SCISSORS.webp",
    "Simurgh's Sovereign": "https://app.chickensaga.com/images/crafting/SIMURGH%27S%20SOVEREIGN.webp",
    "St. Elmo's Fire": "https://app.chickensaga.com/images/crafting/ST%20ELMO%27S%20FIRE.webp",
    "Cocktail's Beak": "https://app.chickensaga.com/images/crafting/DIP%27S%20BEAK.webp",
    "Pos2 Pellet": "https://app.chickensaga.com/images/crafting/POS2%27S%20PELLET.webp",
    "Fetzzz Feet": "https://app.chickensaga.com/images/crafting/FETZZZ%20FEET.webp",
    "Vananderen's Vitality": "https://app.chickensaga.com/images/crafting/VANANDEREN%27S%20VITALITY.webp",
    "Pinong's Bird": "https://app.chickensaga.com/images/crafting/PINONG%27S%20BIRD.webp",
    "Ouchie's Ornament": "https://app.chickensaga.com/images/crafting/OUCHIE%27S%20ORNAMENT.webp",
    "Lockedin State": "https://app.chickensaga.com/images/crafting/LOCKEDIN%20STATE.webp",
}

def get_item_image_url(item_name):
    return ITEM_IMAGE_URLS.get(str(item_name or "").strip(), "")

def is_breedable(chicken):
    return (not chicken.get("is_egg")) and str(chicken.get("state") or "").strip().lower() == "normal"


def safe_int(value, default=None):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default

def apply_gene_profile_classification(record):
    row = dict(record or {})
    gene_profile = classify_gene_profile(row)

    row.update({
        "primary_build": gene_profile.get("primary_build"),
        "primary_build_match_count": gene_profile.get("primary_build_match_count"),
        "primary_build_match_total": gene_profile.get("primary_build_match_total"),
        "primary_build_matched_slots": gene_profile.get("primary_build_matched_slots") or [],
        "primary_build_missing_slots": gene_profile.get("primary_build_missing_slots") or [],
        "primary_build_evaluations": gene_profile.get("primary_build_evaluations") or {},
        "recessive_build": gene_profile.get("recessive_build"),
        "recessive_build_match_count": gene_profile.get("recessive_build_match_count"),
        "recessive_build_match_total": gene_profile.get("recessive_build_match_total"),
        "recessive_build_matched_slots": gene_profile.get("recessive_build_matched_slots") or [],
        "recessive_build_missing_slots": gene_profile.get("recessive_build_missing_slots") or [],
        "recessive_build_repeat_bonus": gene_profile.get("recessive_build_repeat_bonus", 0) or 0,
        "recessive_build_evaluations": gene_profile.get("recessive_build_evaluations") or {},
        "ultimate_type": gene_profile.get("ultimate_type"),
    })

    return row

def require_authorized_wallet(wallet):
    wallet = (wallet or "").strip().lower()
    return wallet and is_authorized_wallet(wallet)


def wallet_needs_daily_refresh(wallet):
    wallet = (wallet or "").strip().lower()
    if not wallet:
        return False

    last_synced_at = get_wallet_last_synced_at(wallet)
    if not last_synced_at:
        return True

    return last_synced_at.astimezone(timezone.utc).date() < datetime.now(timezone.utc).date()


def sync_wallet_data(wallet):
    ensure_static_cache_tables_loaded()

    raw_items = fetch_all_owned_chickens(wallet, CONTRACTS)
    parsed_records = [parse_chicken_record(wallet, item) for item in raw_items]

    current_token_ids = [
        str(row.get("token_id") or "").strip()
        for row in parsed_records
        if str(row.get("token_id") or "").strip()
    ]
    delete_wallet_chickens_not_in_tokens(wallet, current_token_ids)

    static_lookup = get_static_chickens_by_token_ids([row.get("token_id") for row in parsed_records])

    for record in parsed_records:
        token_id = str(record.get("token_id") or "").strip()
        merge_static_chicken_cache(record, static_lookup.get(token_id))

        record = apply_gene_profile_classification(record)
        upsert_chicken(record)

    chickens = get_chickens_by_wallet(wallet)

    preload_cached_family_roots_for_wallet(
        chickens=chickens,
        wallet_address=wallet,
    )

    initialize_simple_family_roots_for_wallet(
        chickens=chickens,
        wallet_address=wallet,
        contract_addresses=CONTRACTS,
    )

    upsert_wallet_last_synced_at(wallet)
    clear_breeding_planner_for_wallet(wallet)

    return get_chickens_by_wallet(wallet)


def get_wallet_chickens(wallet, ensure_loaded=False, force_refresh=False):
    if wallet:
        clear_stale_family_root_summaries(wallet, max_age_hours=24)

    chickens = get_chickens_by_wallet(wallet)
    should_refresh = False

    if ensure_loaded:
        should_refresh = force_refresh or not chickens or wallet_needs_daily_refresh(wallet)

    if should_refresh:
        chickens = sync_wallet_data(wallet)
        if wallet:
            clear_stale_family_root_summaries(wallet, max_age_hours=24)
            chickens = get_chickens_by_wallet(wallet)

    return chickens


def needs_gene_enrichment(chicken):
    return is_breedable(chicken) and needs_recessive_enrichment(chicken)


def enrich_missing_gene_data_in_batches(chickens, wallet, page_key, batch_size=5, prioritized_token_id=None):
    missing = [row for row in (chickens or []) if needs_gene_enrichment(row)]

    cursor_key = f"{page_key}_cursor_{wallet}"

    if not missing:
        session[cursor_key] = 0
        return {
            "loaded": 0,
            "remaining": 0,
        }

    prioritized = []
    remaining = list(missing)

    if prioritized_token_id:
        prioritized = [
            row for row in missing
            if str(row.get("token_id") or "") == str(prioritized_token_id)
        ]
        remaining = [
            row for row in missing
            if str(row.get("token_id") or "") != str(prioritized_token_id)
        ]

    cursor = safe_int(session.get(cursor_key), 0) or 0

    if remaining:
        if cursor >= len(remaining):
            cursor = 0
        rotated = remaining[cursor:] + remaining[:cursor]
    else:
        rotated = []

    selected_batch = prioritized[:1]
    remaining_slots = max(0, batch_size - len(selected_batch))
    selected_batch.extend(rotated[:remaining_slots])

    if not selected_batch:
        remaining_after = len(missing)
        session[cursor_key] = 0
        return {
            "loaded": 0,
            "remaining": remaining_after,
        }

    enriched = enrich_chicken_records(selected_batch)

    for chicken in enriched:
        refreshed = apply_gene_profile_classification(chicken)
        upsert_chicken(refreshed)
        
    if remaining:
        next_cursor = cursor + remaining_slots
        if next_cursor >= len(remaining):
            next_cursor = 0
        session[cursor_key] = next_cursor
    else:
        session[cursor_key] = 0

    refreshed = get_chickens_by_wallet(wallet)
    remaining_after = sum(1 for row in refreshed if needs_gene_enrichment(row))

    return {
        "loaded": len(selected_batch),
        "remaining": remaining_after,
    }


CHICKEN_IMAGE_URL_TEMPLATE = "https://chicken-api-ivory.vercel.app/api/image/{token_id}.png"



def get_chicken_image_url(chicken):
    token_id = str(chicken.get("token_id") or "").strip()
    if not token_id:
        return ""
    existing = str(chicken.get("image") or "").strip()
    if existing:
        return existing
    return CHICKEN_IMAGE_URL_TEMPLATE.format(token_id=token_id)


def enrich_chicken_media(chicken):
    row = dict(chicken)
    row["image"] = get_chicken_image_url(row)
    return row

def get_breeding_planner_session_key(wallet):
    wallet = (wallet or "").strip().lower()
    return f"breeding_planner_queue:{wallet}"

def clear_breeding_planner_for_wallet(wallet):
    wallet = (wallet or "").strip().lower()
    if not wallet:
        return

    session.pop(get_breeding_planner_session_key(wallet), None)
    session.modified = True

def build_planner_pair_key(left_token_id, right_token_id):
    left = str(left_token_id or "").strip()
    right = str(right_token_id or "").strip()
    pair = sorted([left, right])
    return f"{pair[0]}::{pair[1]}" if len(pair) == 2 else ""


def get_breeding_planner_queue(wallet):
    queue = session.get(get_breeding_planner_session_key(wallet)) or []
    if not isinstance(queue, list):
        return []
    cleaned = []
    for row in queue:
        if not isinstance(row, dict):
            continue
        pair_key = str(row.get("pair_key") or "").strip()
        left = row.get("left") or {}
        right = row.get("right") or {}
        if not pair_key or not left.get("token_id") or not right.get("token_id"):
            continue
        cleaned.append(row)
    return cleaned


def save_breeding_planner_queue(wallet, queue_rows):
    session[get_breeding_planner_session_key(wallet)] = queue_rows
    session.modified = True


def get_breeding_planner_token_ids(wallet):
    token_ids = set()
    for row in get_breeding_planner_queue(wallet):
        left = row.get("left") or {}
        right = row.get("right") or {}
        if left.get("token_id"):
            token_ids.add(str(left.get("token_id")))
        if right.get("token_id"):
            token_ids.add(str(right.get("token_id")))
    return token_ids

def filter_out_planner_tokens(rows, wallet):
    queued_ids = get_breeding_planner_token_ids(wallet)
    if not queued_ids:
        return list(rows or [])
    return [row for row in (rows or []) if str(row.get("token_id") or "") not in queued_ids]


def planner_pair_exists(wallet, left_token_id, right_token_id):
    pair_key = build_planner_pair_key(left_token_id, right_token_id)
    return any(
        str(row.get("pair_key") or "") == pair_key
        for row in get_breeding_planner_queue(wallet)
    )


def build_planner_queue_row(mode, left, right, left_item=None, right_item=None, pair_quality="", build_type=""):
    left = enrich_chicken_media(dict(left or {}))
    right = enrich_chicken_media(dict(right or {}))
    left_token_id = str(left.get("token_id") or "").strip()
    right_token_id = str(right.get("token_id") or "").strip()
    pair_key = build_planner_pair_key(left_token_id, right_token_id)
    mode_text = str(mode or "").strip().title()
    build_text = str(build_type or "").strip().title()

    left_summary = left.get("weakest_stat_display") or left.get("build_match_display") or left.get("ultimate_build_match_display") or ""
    right_summary = right.get("weakest_stat_display") or right.get("build_match_display") or right.get("ultimate_build_match_display") or ""

    if str(mode or "").strip().lower() == "ip":
        weakest_info = get_weakest_ip_stat_info(left)
        if weakest_info.get("display"):
            left_summary = weakest_info["display"]
        if weakest_info.get("name"):
            right_summary = f"{weakest_info['label']}: {get_effective_ip_stat(right, weakest_info['name'])}"

    elif str(mode or "").strip().lower() == "gene":
        current_build_type = str(build_type or "").strip().lower()
        build_label = str(build_type or "").strip().title()

        if current_build_type:
            left_gene = enrich_gene_display(left, current_build_type)
            right_gene = enrich_gene_display(right, current_build_type)

            if left_gene.get("build_match_display"):
                left_summary = f"{left_gene['build_match_display']} ({build_label})"
            if right_gene.get("build_match_display"):
                right_summary = f"{right_gene['build_match_display']} ({build_label})"

    elif str(mode or "").strip().lower() == "ultimate":
        left_ultimate = enrich_ultimate_display(left)
        right_ultimate = enrich_ultimate_display(right)

        left_build = str(left_ultimate.get("ultimate_build_display") or "").strip().title()
        right_build = str(right_ultimate.get("ultimate_build_display") or "").strip().title()

        if left_ultimate.get("ultimate_build_match_display"):
            left_summary = (
                f"{left_ultimate['ultimate_build_match_display']} ({left_build})"
                if left_build else left_ultimate["ultimate_build_match_display"]
            )

        if right_ultimate.get("ultimate_build_match_display"):
            right_summary = (
                f"{right_ultimate['ultimate_build_match_display']} ({right_build})"
                if right_build else right_ultimate["ultimate_build_match_display"]
            )

    
    return {
        "pair_key": pair_key,
        "mode": str(mode or "").strip().lower(),
        "mode_label": f"{mode_text} Breeding" if mode_text else "Breeding",
        "build_type": str(build_type or "").strip().lower(),
        "build_label": build_text,
        "pair_quality": str(pair_quality or "").strip(),
        "left": {
            "token_id": left_token_id,
            "image": left.get("image") or "",
            "label": f"#{left_token_id}" if left_token_id else "",
            "ip": left.get("ip"),
            "breed_count": left.get("breed_count"),
            "generation_text": left.get("generation_text") or "",
            "ninuno": left.get("ownership_percent") or 0,
            "summary": left_summary,
        },
        "right": {
            "token_id": right_token_id,
            "image": right.get("image") or "",
            "label": f"#{right_token_id}" if right_token_id else "",
            "ip": right.get("ip"),
            "breed_count": right.get("breed_count"),
            "generation_text": right.get("generation_text") or "",
            "ninuno": right.get("ownership_percent") or 0,
            "summary": right_summary,
        },
        "left_item": left_item or None,
        "right_item": right_item or None,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }


def build_planner_summary(queue_rows):
    queue_rows = list(queue_rows or [])
    count = len(queue_rows)
    return {
        "count": count,
        "label": "Breeding Planner",
        "note": (
            "Queued pairs are removed from available matching pools until you remove them or refresh the wallet."
            if count else
            "You can manually review a match or use Auto Match, then add the pair to the breeding planner."
        ),
        "count_label": "Queued Pairs",
    }



def build_pair_quality_hint(raw_value):
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        value = 0
    if value >= 5:
        return "Excellent match"
    if value >= 4:
        return "Strong match"
    if value >= 2:
        return "Good match"
    return "Situational"


def build_ultimate_pair_quality(row):
    row = row or {}

    left = row.get("left") or {}
    right = row.get("right") or {}
    candidate = row.get("candidate") or {}

    if not right and candidate:
        right = candidate

    build_name = str(
        row.get("selected_build")
        or row.get("build_type")
        or left.get("primary_build")
        or right.get("primary_build")
        or ""
    ).strip().lower()

    left_item = row.get("left_item")
    right_item = row.get("right_item")

    return build_ultimate_pair_quality_from_items(
        left=left,
        right=right,
        build_name=build_name,
        left_item=left_item,
        right_item=right_item,
    )

def export_breeding_planner_excel(queue_rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Breeding Planner"
    headers = [
        "Mode", "Build", "Pair Quality",
        "Left Chicken", "Left IP", "Left Breed Count", "Left Ninuno", "Left Item",
        "Right Chicken", "Right IP", "Right Breed Count", "Right Ninuno", "Right Item",
    ]
    sheet.append(headers)
    for row in queue_rows:
        left = row.get("left") or {}
        right = row.get("right") or {}
        left_item = row.get("left_item") or {}
        right_item = row.get("right_item") or {}
        sheet.append([
            row.get("mode_label") or "",
            row.get("build_label") or "",
            row.get("pair_quality") or "",
            left.get("label") or "", left.get("ip") or "", left.get("breed_count") if left.get("breed_count") is not None else "", left.get("ninuno") or 0, left_item.get("name") or "",
            right.get("label") or "", right.get("ip") or "", right.get("breed_count") if right.get("breed_count") is not None else "", right.get("ninuno") or 0, right_item.get("name") or "",
        ])
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 4, 24)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

def format_wallet_short(wallet):
    wallet = str(wallet or "").strip()
    if len(wallet) <= 14:
        return wallet
    return f"{wallet[:6]}...{wallet[-4:]}"


def format_wallet_updated_at(dt_value):
    if not dt_value:
        return ""

    try:
        local_dt = dt_value.astimezone()
    except Exception:
        local_dt = dt_value

    month = local_dt.strftime("%b")
    day = local_dt.day
    year = local_dt.year
    hour = local_dt.strftime("%I").lstrip("0") or "0"
    minute = local_dt.strftime("%M")
    am_pm = local_dt.strftime("%p")
    return f"{month} {day}, {year} {hour}:{minute} {am_pm}"


def build_count_distribution_rows(pairs):
    pairs = list(pairs or [])
    total = sum(max(0, int(item.get("count", 0) or 0)) for item in pairs)
    rows = []

    for item in pairs:
        count = max(0, int(item.get("count", 0) or 0))
        percent = round((count / total) * 100, 1) if total else 0
        rows.append({
            "key": str(item.get("key") or "").strip().lower(),
            "label": str(item.get("label") or "").strip(),
            "count": count,
            "percent": percent,
        })

    return rows


def build_available_chickens_dashboard(chickens, breedable_chickens):
    chickens = list(chickens or [])
    breedable_chickens = list(breedable_chickens or [])

    total_count = len(chickens)
    breedable_count = len(breedable_chickens)
    egg_count = sum(
        1
        for row in chickens
        if bool(row.get("is_egg")) or str(row.get("type") or "").strip().lower() == "egg"
    )
    unavailable_count = max(0, total_count - breedable_count - egg_count)

    composition = build_count_distribution_rows([
        {"key": "breedable", "label": "Breedable", "count": breedable_count},
        {"key": "eggs", "label": "Eggs", "count": egg_count},
        {"key": "unavailable", "label": "Unavailable", "count": unavailable_count},
    ])

    generation_counts = {}
    for row in breedable_chickens:
        label = str(row.get("generation_text") or "").strip() or "Unknown"
        sort_value = safe_int(row.get("generation_num"), 999999)
        current = generation_counts.get(label, {"count": 0, "sort_value": sort_value})
        current["count"] += 1
        current["sort_value"] = min(current["sort_value"], sort_value)
        generation_counts[label] = current

    generation_distribution = build_count_distribution_rows([
        {
            "key": label.lower(),
            "label": label,
            "count": data["count"],
        }
        for label, data in sorted(
            generation_counts.items(),
            key=lambda item: (item[1]["sort_value"], item[0].lower())
        )
    ])

    breed_count_counts = {}
    for row in breedable_chickens:
        breed_count = safe_int(row.get("breed_count"), 0)
        label = str(breed_count if breed_count is not None else 0)
        breed_count_counts[label] = breed_count_counts.get(label, 0) + 1

    breed_count_distribution = build_count_distribution_rows([
        {
            "key": f"breed-{label}",
            "label": label,
            "count": count,
        }
        for label, count in sorted(
            breed_count_counts.items(),
            key=lambda item: safe_int(item[0], 999999)
        )
    ])

    ninuno_counts = {
        "complete_100": 0,
        "above_0": 0,
        "partial": 0,
        "recalculate": 0,
    }

    for row in breedable_chickens:
        ownership = float(row.get("ownership_percent") or 0)
        total_root_count = safe_int(row.get("total_root_count"), 0) or 0
        is_complete = bool(row.get("is_complete"))

        if not total_root_count:
            ninuno_counts["recalculate"] += 1
        elif not is_complete:
            ninuno_counts["partial"] += 1
        elif ownership == 100.0:
            ninuno_counts["complete_100"] += 1
        elif ownership > 0:
            ninuno_counts["above_0"] += 1
        else:
            ninuno_counts["above_0"] += 1

    ninuno_distribution = build_count_distribution_rows([
        {"key": "complete_100", "label": "100% Ninuno", "count": ninuno_counts["complete_100"]},
        {"key": "above_0", "label": "Above 0%", "count": ninuno_counts["above_0"]},
        {"key": "partial", "label": "Partial", "count": ninuno_counts["partial"]},
        {"key": "recalculate", "label": "Recalculate", "count": ninuno_counts["recalculate"]},
    ])

    highest_ip_row = None
    if breedable_chickens:
        highest_ip_row = max(
            breedable_chickens,
            key=lambda row: (
                safe_int(row.get("ip"), -1) or -1,
                -(safe_int(row.get("breed_count"), 999999) or 999999),
                -(safe_int(row.get("generation_num"), 999999) or 999999),
            )
        )

    lowest_breed_count_row = None
    if breedable_chickens:
        lowest_breed_count_row = min(
            breedable_chickens,
            key=lambda row: (
                safe_int(row.get("breed_count"), 999999) or 999999,
                -(safe_int(row.get("ip"), -1) or -1),
                safe_int(row.get("generation_num"), 999999) or 999999,
            )
        )

    return {
        "composition": composition,
        "generation_distribution": generation_distribution,
        "breed_count_distribution": breed_count_distribution,
        "ninuno_distribution": ninuno_distribution,
        "highest_ip_row": enrich_chicken_media(highest_ip_row) if highest_ip_row else None,
        "lowest_breed_count_row": enrich_chicken_media(lowest_breed_count_row) if lowest_breed_count_row else None,
    }

def build_wallet_summary(wallet, chickens, access_expiry=None):
    chickens = list(chickens or [])

    total_count = len(chickens)
    breedable_count = sum(1 for row in chickens if is_breedable(row))
    egg_count = sum(
        1
        for row in chickens
        if bool(row.get("is_egg")) or str(row.get("type") or "").strip().lower() == "egg"
    )
    unavailable_count = max(0, total_count - breedable_count - egg_count)
    planner_count = len(get_breeding_planner_queue(wallet))

    last_synced_at = get_wallet_last_synced_at(wallet)
    updated_display = format_wallet_updated_at(last_synced_at)

    return {
        "wallet_full": wallet,
        "wallet_short": format_wallet_short(wallet),
        "total_count": total_count,
        "breedable_count": breedable_count,
        "egg_count": egg_count,
        "unavailable_count": unavailable_count,
        "planner_count": planner_count,
        "access_expiry": access_expiry or "",
        "updated_display": updated_display,
    }

@app.context_processor
def inject_breeding_item_helpers():
    return {
        "recommend_ip_item": recommend_ip_item,
        "recommend_gene_item": recommend_gene_item,
        "get_ip_item_candidates": get_ip_item_candidates,
        "get_gene_item_candidates": get_gene_item_candidates,
        "get_ultimate_item_candidates": get_ultimate_item_candidates,
        "resolve_pair_item_recommendations": resolve_pair_item_recommendations,
        "resolve_ultimate_pair_item_recommendations": resolve_ultimate_pair_item_recommendations,
        "get_item_image_url": get_item_image_url,
        "build_ip_pair_quality": build_ip_pair_quality,
        "build_gene_pair_quality": build_gene_pair_quality,
        "build_ultimate_pair_quality": build_ultimate_pair_quality,
        "build_prefers_instinct": build_prefers_instinct,
        "planner_pair_exists": lambda wallet, left_token_id, right_token_id: planner_pair_exists(wallet, left_token_id, right_token_id),
    }


def enrich_gene_display(chicken, build_type):
    row = dict(chicken)

    row["build_source_display"] = ""
    row["build_match_display"] = ""
    row["gene_sort_source_rank"] = 9
    row["gene_sort_match_count"] = 0
    row["gene_sort_match_total"] = 0
    row["gene_effective_source"] = ""

    if not build_type:
        return row

    target_info = get_gene_build_target_info(row, build_type)
    row["build_source_display"] = target_info["display_source"]
    row["build_match_display"] = target_info["display_match"]
    row["gene_sort_source_rank"] = target_info["sort_source_rank"]
    row["gene_sort_match_count"] = target_info["sort_match_count"]
    row["gene_sort_match_total"] = target_info["sort_match_total"]
    row["gene_effective_source"] = target_info["source"]

    return enrich_chicken_media(row)

def enrich_gene_available_display(chicken):
    row = enrich_chicken_media(dict(chicken or {}))

    best_info = get_best_available_gene_build_info(row)

    row["build_type"] = best_info.get("build_key") or ""
    row["build_label"] = best_info.get("build_label") or ""
    row["build_match_display"] = best_info.get("build_count_display") or ""
    row["build_match_count"] = best_info.get("sort_match_count", 0) or 0
    row["build_match_total"] = best_info.get("sort_match_total", 0) or 0
    row["gene_sort_source_rank"] = best_info.get("sort_source_rank", 99)
    row["gene_sort_match_count"] = best_info.get("sort_match_count", 0) or 0
    row["gene_sort_match_total"] = best_info.get("sort_match_total", 0) or 0

    build_type = row["build_type"]
    if build_type:
        target_info = get_gene_build_target_info(row, build_type)
        row["build_source_display"] = target_info.get("display_source") or ""
        row["gene_effective_source"] = target_info.get("source") or ""
    else:
        row["build_source_display"] = ""
        row["gene_effective_source"] = ""

    return row

def chicken_matches_gene_build(chicken, build_type):
    if not build_type:
        return False

    return get_gene_build_target_info(chicken, build_type)["eligible"]


def build_ultimate_available_empty_state():
    return make_empty_state(
        "No candidates ready",
        "No ultimate-eligible chickens found.",
        "Load more breedable chickens or review your wallet inventory before returning to Ultimate matching.",
    )


def build_ultimate_match_empty_state(auto_match=False, ninuno_mode="all", breed_diff=None):
    if breed_diff is not None and breed_diff <= 1:
        return make_empty_state(
            "Breed count too tight",
            "No ultimate pair passed the current breed-count difference rule.",
            "Increase the allowed breed-count difference to widen the partner pool.",
        )
    if str(ninuno_mode or "all") != "all":
        return make_empty_state(
            "Ninuno filter active",
            "No ultimate pair survived the current Ninuno filter.",
            "Relax the Ninuno filter to include more candidates.",
        )
    if auto_match:
        return make_empty_state(
            "No valid auto-match",
            "Auto Match could not find a usable ultimate pair.",
            "Try wider popup filters or choose a chicken manually.",
        )
    return make_empty_state(
        "No matches found",
        "No valid ultimate pair was found for the selected chicken.",
        "Try another chicken or review a larger available pool.",
    )

def make_empty_state(kicker, title, body):
    return {
        "kicker": kicker,
        "title": title,
        "body": body,
    }


def build_gene_available_empty_state(build_type, ninuno_100_only=False):
    if not build_type:
        return make_empty_state(
            "Choose a build",
            "Select a build to see available chickens.",
            "Pick your target build first to load compatible gene candidates.",
        )
    if ninuno_100_only:
        return make_empty_state(
            "Build plus Ninuno filter",
            "No chickens matched this build under the 100% Ninuno filter.",
            "Turn off the 100% Ninuno filter or try a different build.",
        )
    return make_empty_state(
        "Build has no candidates",
        "No chickens were found for this build.",
        "Try another build or refresh the wallet if you expect more candidates.",
    )


def build_gene_match_empty_state(build_type, ninuno_100_only=False, auto_match=False, same_instinct=False, min_build_count=None):
    if min_build_count is not None and min_build_count >= 4:
        return make_empty_state(
            "Build threshold too high",
            "No gene pair passed the current build-match requirement.",
            "Lower the minimum build-match count to widen the pool.",
        )
    if same_instinct:
        return make_empty_state(
            "Same instinct required",
            "No gene pair survived the same-instinct filter.",
            "Turn off Same Instinct to allow more compatible partners.",
        )
    if ninuno_100_only:
        return make_empty_state(
            "100% Ninuno only",
            "No gene pair survived the Ninuno filter.",
            "Turn off the 100% Ninuno filter to allow more possible partners.",
        )
    if auto_match:
        return make_empty_state(
            "No valid auto-match",
            "Auto Match could not find a usable gene pair.",
            "Try a different build, wider popup filters, or choose a chicken manually.",
        )
    build_label = str(build_type or "selected").title()
    return make_empty_state(
        "No matches found",
        f"No valid {build_label} gene pair was found for the selected chicken.",
        "Try another chicken or a different build target.",
    )


def enrich_ultimate_display(chicken):
    row = dict(chicken)
    best_info = get_best_available_gene_build_info(row)

    primary_build = normalize_ultimate_build_value(row.get("primary_build"))
    if not primary_build:
        primary_build = (
            str(row.get("gene_build_key") or "").strip().lower()
            or str(best_info.get("build_key") or "").strip().lower()
        )
        row["primary_build"] = primary_build

    primary_count = safe_int(row.get("primary_build_match_count"), 0) or 0
    primary_total = safe_int(row.get("primary_build_match_total"), 0) or 0
    if not primary_count and safe_int(row.get("gene_build_match_count"), 0):
        row["primary_build_match_count"] = safe_int(row.get("gene_build_match_count"), 0) or 0
        primary_count = row["primary_build_match_count"]
    elif not primary_count and safe_int(best_info.get("sort_match_count"), 0):
        row["primary_build_match_count"] = safe_int(best_info.get("sort_match_count"), 0) or 0
        primary_count = row["primary_build_match_count"]
    if not primary_total and safe_int(row.get("gene_build_match_total"), 0):
        row["primary_build_match_total"] = safe_int(row.get("gene_build_match_total"), 0) or 0
        primary_total = row["primary_build_match_total"]
    elif not primary_total and safe_int(best_info.get("sort_match_total"), 0):
        row["primary_build_match_total"] = safe_int(best_info.get("sort_match_total"), 0) or 0
        primary_total = row["primary_build_match_total"]

    row["ultimate_type_display"] = get_ultimate_type_display(row)
    row["ultimate_build_display"] = get_ultimate_build_display(row)
    row["ultimate_build_match_display"] = f"{primary_count}/{primary_total}" if primary_total else ""

    return enrich_chicken_media(row)

def pick_best_ultimate_auto_match(breedable_chickens, include_lower_values=False):
    return service_pick_best_ultimate_auto_match(
        breedable_chickens,
        include_lower_values=include_lower_values,
    )

def parse_build_match_count(value):
    raw = str(value or "").strip()
    if not raw:
        return 0

    best = 0
    for part in raw.split("+"):
        part = part.strip()
        if "/" in part:
            left, _, _ = part.partition("/")
            best = max(best, safe_int(left, 0) or 0)
        else:
            best = max(best, safe_int(part, 0) or 0)
    return best

def build_ultimate_available_auto_candidates(breedable_chickens, breed_diff=None, ninuno_mode="all", include_lower_values=False):
    return service_build_ultimate_available_auto_candidates(
        breedable_chickens=breedable_chickens,
        breed_diff=breed_diff,
        ninuno_mode=ninuno_mode,
        include_lower_values=include_lower_values,
    )


def pick_multi_pairs_from_candidates(pair_candidates, target_count, mode=""):
    pair_candidates = list(pair_candidates or [])
    mode_key = str(mode or "").strip().lower()

    if not mode_key:
        if any("gene_pair_metrics" in row or "selected_eval" in row for row in pair_candidates):
            mode_key = "gene"
        elif any("ultimate_build_metrics" in row or "pair_quality" in row for row in pair_candidates):
            mode_key = "ultimate"

    def safe_pair_ranking(row):
        return tuple(row.get("ranking") or ())

    def safe_side_int(row, side, keys):
        side_row = (row.get(side) or {})
        for key in keys:
            direct_value = safe_int(row.get(f"{side}_{key}"))
            if direct_value is not None:
                return direct_value
            value = safe_int(side_row.get(key))
            if value is not None:
                return value
        return 0

    def sort_key(row):
        if mode_key == "gene":
            left_value = safe_side_int(row, "left", ("build_match_count",))
            right_value = safe_side_int(row, "right", ("build_match_count",))
            return (
                -max(left_value, right_value),
                -min(left_value, right_value),
                -(left_value + right_value),
                safe_pair_ranking(row),
            )

        if mode_key == "ultimate":
            return safe_pair_ranking(row)

        return safe_pair_ranking(row)

    used = set()
    results = []
    target_count = max(0, safe_int(target_count, 0) or 0)
    ordered_candidates = sorted(pair_candidates, key=sort_key)
    for row in ordered_candidates:
        left_id = str((row.get("left") or {}).get("token_id") or "")
        right_id = str((row.get("right") or {}).get("token_id") or "")
        if not left_id or not right_id or left_id in used or right_id in used:
            continue
        results.append(row)
        used.add(left_id)
        used.add(right_id)
        if len(results) >= target_count:
            break
    return results

def sort_key_text(value):
    return str(value or "").strip().lower()


def sort_key_int(value, default=0):
    parsed = safe_int(value, default)
    return parsed if parsed is not None else default


def get_gene_build_source_rank(value):
    source = str(value or "").strip().lower()
    if source == "primary":
        return 0
    if source in {"primary + recessive", "mixed"}:
        return 1
    if source == "recessive":
        return 2
    return 9

def get_ip_difference(chicken_a, chicken_b):
    ip_a = safe_int((chicken_a or {}).get("ip"))
    ip_b = safe_int((chicken_b or {}).get("ip"))

    if ip_a is None or ip_b is None:
        return None

    return abs(ip_a - ip_b)

register_core_routes(app, {
    "build_available_chickens_dashboard": build_available_chickens_dashboard,
    "build_wallet_summary": build_wallet_summary,
    "clear_owner_admin_failures": clear_owner_admin_failures,
    "enrich_chicken_media": enrich_chicken_media,
    "format_wallet_access_rows": format_wallet_access_rows,
    "get_best_available_gene_build_info": get_best_available_gene_build_info,
    "get_owner_admin_password": lambda: OWNER_ADMIN_PASSWORD,
    "get_wallet_access_expiry_display": get_wallet_access_expiry_display,
    "get_wallet_access_rows": get_wallet_access_rows,
    "get_wallet_chickens": get_wallet_chickens,
    "grant_manual_access": grant_manual_access,
    "has_wallet_access": has_wallet_access,
    "is_authorized_wallet": is_authorized_wallet,
    "is_breedable": is_breedable,
    "is_owner_admin_locked": is_owner_admin_locked,
    "is_valid_wallet": is_valid_wallet,
    "owner_password_is_valid": owner_password_is_valid,
    "owner_whitelist_route": OWNER_WHITELIST_ROUTE,
    "register_owner_admin_failure": register_owner_admin_failure,
    "require_authorized_wallet": require_authorized_wallet,
    "safe_int": safe_int,
    "set_authorized_wallet": set_authorized_wallet,
    "static_export_db_path": STATIC_EXPORT_DB_PATH,
    "sync_static_export_tables_to_main_db": sync_static_export_tables_to_main_db,
    "sync_wallet_data": sync_wallet_data,
})

register_planner_routes(app, {
    "build_apex_breeder_bookmarklet_code": build_apex_breeder_bookmarklet_code,
    "build_bookmarklet_inventory_name_lookup": build_bookmarklet_inventory_name_lookup,
    "build_per_pair_item_status": build_per_pair_item_status,
    "build_planner_queue_row": build_planner_queue_row,
    "build_planner_summary": build_planner_summary,
    "build_wallet_inventory_lookup": build_wallet_inventory_lookup,
    "build_wallet_planner_item_requirements_summary": build_wallet_planner_item_requirements_summary,
    "build_wallet_summary": build_wallet_summary,
    "enrich_chicken_media": enrich_chicken_media,
    "export_breeding_planner_excel": export_breeding_planner_excel,
    "get_breeding_planner_queue": get_breeding_planner_queue,
    "get_wallet_access_expiry_display": get_wallet_access_expiry_display,
    "get_wallet_chickens": get_wallet_chickens,
    "planner_pair_exists": planner_pair_exists,
    "require_authorized_wallet": require_authorized_wallet,
    "save_breeding_planner_queue": save_breeding_planner_queue,
})

register_match_routes(app, {
    "build_gene_active_filters": build_gene_active_filters,
    "build_gene_available_auto_candidates_same_build": build_gene_available_auto_candidates_same_build,
    "build_gene_available_filter_options": build_gene_available_filter_options,
    "build_gene_match_empty_state": build_gene_match_empty_state,
    "build_gene_potential_matches_strict": build_gene_potential_matches_strict,
    "build_ip_active_filters": build_ip_active_filters,
    "build_ip_available_filter_options": build_ip_available_filter_options,
    "build_ip_multi_matches": build_ip_multi_matches,
    "build_planner_summary": build_planner_summary,
    "build_ultimate_active_filters": build_ultimate_active_filters,
    "build_ultimate_available_auto_candidates": build_ultimate_available_auto_candidates,
    "build_ultimate_available_empty_state": build_ultimate_available_empty_state,
    "build_ultimate_available_filter_options": build_ultimate_available_filter_options,
    "build_ultimate_match_empty_state": build_ultimate_match_empty_state,
    "build_wallet_summary": build_wallet_summary,
    "chicken_matches_gene_available_filters": chicken_matches_gene_available_filters,
    "chicken_matches_ip_available_filters": chicken_matches_ip_available_filters,
    "chicken_matches_ultimate_available_filters": chicken_matches_ultimate_available_filters,
    "chicken_passes_auto_ninuno_filter": chicken_passes_auto_ninuno_filter,
    "complete_ninuno_via_lineage_with_resume": complete_ninuno_via_lineage_with_resume,
    "CONTRACTS": CONTRACTS,
    "enrich_chicken_media": enrich_chicken_media,
    "enrich_gene_available_chicken_row": enrich_gene_available_chicken_row,
    "enrich_ip_available_chicken_row": enrich_ip_available_chicken_row,
    "enrich_missing_gene_data_in_batches": enrich_missing_gene_data_in_batches,
    "enrich_missing_recessive_data_in_batches": enrich_missing_recessive_data_in_batches,
    "enrich_ultimate_available_chicken_row": enrich_ultimate_available_chicken_row,
    "enrich_ultimate_display": enrich_ultimate_display,
    "filter_and_sort_ultimate_candidates": filter_and_sort_ultimate_candidates,
    "filter_out_planner_tokens": filter_out_planner_tokens,
    "find_potential_matches": find_potential_matches,
    "GENE_BUILD_ORDER": GENE_BUILD_ORDER,
    "get_best_available_gene_build_info": get_best_available_gene_build_info,
    "get_breeding_planner_queue": get_breeding_planner_queue,
    "get_chickens_by_wallet": get_chickens_by_wallet,
    "get_effective_ip_stat": get_effective_ip_stat,
    "get_featured_market_feed": get_featured_market_feed,
    "get_ultimate_build_display": get_ultimate_build_display,
    "get_ultimate_type_display": get_ultimate_type_display,
    "get_wallet_access_expiry_display": get_wallet_access_expiry_display,
    "get_wallet_chickens": get_wallet_chickens,
    "get_weakest_ip_stat_info": get_weakest_ip_stat_info,
    "has_active_payment_access_in_db": has_active_payment_access_in_db,
    "is_breedable": is_breedable,
    "is_full_siblings": is_full_siblings,
    "is_generation_gap_allowed": is_generation_gap_allowed,
    "is_parent_offspring": is_parent_offspring,
    "is_ultimate_eligible": is_ultimate_eligible,
    "make_empty_state": make_empty_state,
    "match_settings": MATCH_SETTINGS,
    "normalize_auto_ninuno_filter": normalize_auto_ninuno_filter,
    "normalize_gene_available_build_filter": normalize_gene_available_build_filter,
    "normalize_gene_available_ninuno_filter": normalize_gene_available_ninuno_filter,
    "normalize_gene_available_source_values": normalize_gene_available_source_values,
    "normalize_ip_available_ninuno_filter": normalize_ip_available_ninuno_filter,
    "normalize_ultimate_available_ninuno_filter": normalize_ultimate_available_ninuno_filter,
    "pair_has_usable_ip_items": pair_has_usable_ip_items,
    "parse_gene_csv_query_values": parse_gene_csv_query_values,
    "parse_ip_csv_query_values": parse_csv_query_values,
    "parse_ultimate_csv_query_values": parse_ultimate_csv_query_values,
    "pick_best_gene_auto_match_from_pool": pick_best_gene_auto_match_from_pool,
    "pick_best_ultimate_auto_match": pick_best_ultimate_auto_match,
    "pick_multi_pairs_from_candidates": pick_multi_pairs_from_candidates,
    "refresh_ultimate_primary_builds_if_needed": refresh_ultimate_primary_builds_if_needed,
    "require_authorized_wallet": require_authorized_wallet,
    "safe_int": safe_int,
    "sort_gene_available_chickens": sort_gene_available_chickens,
    "sort_ip_available_chickens": sort_ip_available_chickens,
    "sort_ip_match_rows": sort_ip_match_rows,
    "sort_key_int": sort_key_int,
    "sort_key_text": sort_key_text,
    "sort_ultimate_available_table_chickens": sort_ultimate_available_table_chickens,
    "upsert_chicken": upsert_chicken,
    "upsert_family_root_summary": upsert_family_root_summary,
})

if __name__ == "__main__":
    app.run(debug=FLASK_DEBUG_ENABLED)
